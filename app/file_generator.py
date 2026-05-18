from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet


def generate_docx(content, output_path):
    doc = Document()
    doc.add_heading('VisionRAG Report')
    doc.add_paragraph(content)
    doc.save(output_path)


def generate_pdf(content, output_path):
    pdf = SimpleDocTemplate(output_path)
    styles = getSampleStyleSheet()

    story = [Paragraph(content, styles['BodyText'])]
    pdf.build(story)